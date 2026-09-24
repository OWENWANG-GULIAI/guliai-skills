#!/usr/bin/env python3
"""Build the neutral default Word template used by this Skill when no client template exists."""

import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 99, 112)
LIGHT = "F4F6F9"


def set_font(run, name="Calibri", size=11, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_font(run, size=10.5, color=color, bold=bold)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def main():
    if len(sys.argv) != 2:
        print("Usage: python build_default_word_template.py <output.docx>")
        raise SystemExit(2)
    output = Path(sys.argv[1])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for level, size, color, before, after in [(1, 16, BLUE, 18, 10), (2, 13, BLUE, 12, 6), (3, 12, RGBColor(31, 77, 120), 8, 4)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    subtitle = doc.styles.add_style("Proposal Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(20)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_p.add_run("企业内训咨询方案")
    set_font(header_run, size=9, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("[企业名称 / 或客户名称]")
    set_font(footer_run, size=9, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("[课程名称]")
    set_font(run, size=25, color=INK, bold=True)
    p = doc.add_paragraph(style="Proposal Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("企业内训课程方案")

    meta = doc.add_table(rows=2, cols=2)
    set_table_widths(meta, [3.25, 3.25])
    meta.style = "Table Grid"
    mark_header_row(meta.rows[0])
    fields = [("服务对象", "[客户名称 / 部门]"), ("方案版本", "[日期 / 版本号]"), ("培训形式", "[线上 / 线下 / 混合]"), ("培训时长", "[时长与时间安排]")]
    for cell, (label, value) in zip([c for r in meta.rows for c in r.cells], fields):
        shade(cell, LIGHT)
        set_cell_text(cell, f"{label}：{value}")

    add_heading(doc, "一、执行摘要")
    doc.add_paragraph("[用三至五句话概述客户要解决的问题、推荐方案与预期成果。]")
    add_heading(doc, "二、需求理解")
    for text in ["已确认需求：[客户明确提供的信息]", "课程设计判断：[基于事实形成的建议]", "待确认事项：[仍影响范围或交付的事项]"]:
        doc.add_paragraph(text)
    add_heading(doc, "三、课程定位与培训目标")
    doc.add_paragraph("适用对象：[学员画像]")
    doc.add_paragraph("培训目标：[培训后，学员能够……]")
    doc.add_paragraph("课程设计逻辑：[理解 → 示范 → 练习 → 迁移 → 复盘]")
    add_heading(doc, "四、课程安排")
    agenda = doc.add_table(rows=2, cols=5)
    agenda.style = "Table Grid"
    mark_header_row(agenda.rows[0])
    set_table_widths(agenda, [0.8, 1.25, 2.0, 1.2, 1.25])
    headers = ["时间", "模块", "核心内容", "学习方式", "模块产出"]
    for cell, text in zip(agenda.rows[0].cells, headers):
        shade(cell, "E8EEF5")
        set_cell_text(cell, text, bold=True, color=INK)
    for cell, text in zip(agenda.rows[1].cells, ["[时段]", "[模块名称]", "[内容]", "[方式]", "[产出]"]):
        set_cell_text(cell, text)
    add_heading(doc, "五、交付与实施建议")
    doc.add_paragraph("[课程结束后留下的成果、推荐形式及调整原则。]")
    add_heading(doc, "六、培训前准备与范围边界")
    doc.add_paragraph("客户需配合：[材料、设备、组织安排及脱敏要求。]")
    doc.add_paragraph("范围边界：[暂不包含项、实施依赖与待确认项。]")
    doc.save(output)
    print(f"[OK] Created {output}")


if __name__ == "__main__":
    main()
